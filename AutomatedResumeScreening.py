# resume_screening_nlp/project.py (BERT + Streamlit + Filters + Visuals + API)

import re
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
from pdfminer.high_level import extract_text
from docx import Document
from sentence_transformers import SentenceTransformer, util
from flask import Flask, request, jsonify

# === Load Sentence Transformer Model ===
model = SentenceTransformer('all-MiniLM-L6-v2')

# === Resume Parsing ===
def extract_text_from_pdf(file_obj):
    try:
        return extract_text(file_obj)
    except:
        return ""

def extract_text_from_docx(file_obj):
    try:
        doc = Document(file_obj)
        return " ".join([para.text for para in doc.paragraphs])
    except:
        return ""

def parse_resume(uploaded_file):
    file_name = uploaded_file.name.lower()
    if file_name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif file_name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return ""

# === Text Cleaning ===
def clean_text(text):
    text = re.sub(r"[^a-zA-Z0-9 ]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.lower()

# === Information Extraction ===
def extract_experience(text):
    match = re.findall(r'(\d+)\s*(?:years|yrs)', text)
    if match:
        return max([int(x) for x in match])
    return 0

def extract_education(text):
    education_keywords = ['bachelor', 'btech', 'master', 'mba', 'bca', 'mca', 'phd', 'msc', 'bsc']
    for keyword in education_keywords:
        if keyword in text:
            return keyword.upper()
    return "UNKNOWN"

def extract_marks(text):
    match = re.findall(r'(?:marks|scored|percentage)?\s*(\d{2,3})\s*%?', text)
    if match:
        return max([int(x) for x in match if int(x) <= 100])
    return None

# === Skill Matching ===
def analyze_skills(text, required_skills):
    matched = [skill for skill in required_skills if skill in text]
    missing = [skill for skill in required_skills if skill not in text]
    return matched, missing

# === Resume Ranking ===
def rank_resumes_bert(job_description, resume_texts, min_exp, min_marks, required_degree):
    resume_embeddings = model.encode(resume_texts, convert_to_tensor=True)
    combined_query = f"{job_description}. Minimum {min_exp} years experience. Minimum marks {min_marks}. Degree required: {required_degree}."
    combined_query_embedding = model.encode(combined_query, convert_to_tensor=True)
    scores = util.cos_sim(combined_query_embedding, resume_embeddings)[0]
    return scores.cpu().numpy()

# === Streamlit UI ===
def run_streamlit():
    st.title("📄 Automated Resume Screening with BERT")

    job_title = st.text_input("Enter Job Title/Keyword")
    skill_filter_input = st.text_input("Enter Required Skills (comma-separated, e.g., Python, SQL, ML)")
    required_skills = [s.strip().lower() for s in skill_filter_input.split(",") if s.strip()]
    min_exp = st.slider("Minimum Experience (Years)", 0, 20, 0)
    required_degree = st.selectbox("Required Education (Optional)", ["Any", "BTECH", "MCA", "BCA", "MBA", "PHD", "BSC", "MSC"])
    min_marks = st.slider("Minimum Percentage (Optional)", 0, 100, 0)

    job_desc_path = st.file_uploader("Upload Job Description (TXT/PDF/DOCX)", type=["txt", "pdf", "docx"])
    resumes = st.file_uploader("Upload Resumes (PDF/DOCX)", type=["pdf", "docx"], accept_multiple_files=True)

    if job_desc_path is not None and resumes:
        file_name = job_desc_path.name.lower()
        if file_name.endswith(".pdf"):
            job_description = clean_text(extract_text(job_desc_path))
        elif file_name.endswith(".docx"):
            job_description = clean_text(" ".join([p.text for p in Document(job_desc_path).paragraphs]))
        else:
            job_description = clean_text(job_desc_path.read().decode("utf-8", errors="ignore"))

        resume_texts = [clean_text(parse_resume(resume)) for resume in resumes]

        st.write("✅ Job Description Sample:", job_description[:300])
        st.write("✅ Loaded Resume Text Samples:", [r[:300] for r in resume_texts])

        if not job_description.strip() or not any(resume_texts):
            st.warning("Job description or all resumes are empty or unparseable.")
            return

        scores = rank_resumes_bert(job_description, resume_texts, min_exp, min_marks, required_degree)

        results = pd.DataFrame({
            'Resume File': [resume.name for resume in resumes],
            'Raw Text': resume_texts,
            'Score': scores,
        })

        results['Experience (yrs)'] = [extract_experience(text) for text in resume_texts]
        results['Education'] = [extract_education(text) for text in resume_texts]
        results['Marks (%)'] = [extract_marks(text) for text in resume_texts]

        if required_skills:
            original_count = len(results)
            results = results[results['Raw Text'].apply(lambda text: any(skill in text for skill in required_skills))]
            st.info(f"{len(results)} out of {original_count} resumes matched at least one of the required skills: {', '.join(required_skills)}")

            if not results.empty:
                results['Matched Skills'], results['Missing Skills'] = zip(*results['Raw Text'].apply(lambda text: analyze_skills(text, required_skills)))

        results = results[results['Experience (yrs)'] >= min_exp]

        if required_degree != "Any":
            results = results[results['Education'] == required_degree.upper()]

        results['Marks (%)'] = results['Marks (%)'].fillna(0)
        results = results[results['Marks (%)'] >= min_marks]

        if not results.empty:
            results = results.sort_values(by='Score', ascending=False).reset_index(drop=True)
            results.index += 1  # Start ranking from 1
            columns_to_display = ['Resume File', 'Score', 'Experience (yrs)', 'Education', 'Marks (%)']
            if required_skills:
                columns_to_display += ['Matched Skills', 'Missing Skills']

            st.subheader("📊 Ranked Resumes")
            st.dataframe(results[columns_to_display].rename_axis("Rank").reset_index())

            st.subheader("📈 Similarity Score Distribution")
            fig, ax = plt.subplots()
            ax.hist(results['Score'], bins=10, color='skyblue', edgecolor='black')
            ax.set_xlabel('Similarity Score')
            ax.set_ylabel('Number of Resumes')
            st.pyplot(fig)

            st.download_button("Download Results", data=results[columns_to_display].rename_axis("Rank").reset_index().to_csv(index=False), file_name="filtered_ranked_resumes.csv")
        else:
            st.warning("No resumes matched the filters or scoring criteria.")
    else:
        st.info("Please upload both a Job Description and Resumes.")

# === Optional Flask API ===
app = Flask(__name__)

@app.route('/api/screen', methods=['POST'])
def api_screen():
    data = request.json
    job_desc = clean_text(data.get("job_description", ""))
    resumes = data.get("resumes", [])
    scores = rank_resumes_bert(job_desc, [clean_text(text) for text in resumes], 0, 0, "Any")
    result = [{"resume": f"Resume_{i+1}", "score": float(score)} for i, score in enumerate(scores)]
    result.sort(key=lambda x: x["score"], reverse=True)
    return jsonify(result)

# === Entry Point ===
if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == 'api':
        app.run(debug=True, port=5000)
    else:
        run_streamlit()


# streamlit run AutomatedResumeScreening.py

